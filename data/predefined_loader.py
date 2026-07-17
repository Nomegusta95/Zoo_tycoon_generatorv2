import os
from openpyxl import load_workbook

REQUIRED_PROJECT_COLUMNS = ["name", "animals"]


def _read_rows(path):
    """Yields (excel_row_number, {column_name: cell_value}) for each data
    row - every row dict has every header column as a key (even when the
    cell itself is blank, i.e. None), matching how a pandas row exposes
    every DataFrame column via `in`/`.get()` regardless of blank cells."""
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active

    rows = ws.iter_rows(values_only=True)
    header = [str(c).strip() if c is not None else "" for c in next(rows)]

    for excel_row, values in enumerate(rows, start=2):
        yield excel_row, dict(zip(header, values))


def load_predefined_projects(path):
    if not os.path.exists(path):
        raise FileNotFoundError(
            f"Predefined projects file not found: {path}\n"
            "Check that the file exists at that path."
        )

    try:
        rows = list(_read_rows(path))
    except Exception as e:
        raise ValueError(f"Could not read predefined projects file '{path}': {e}") from e

    columns = list(rows[0][1].keys()) if rows else []

    missing = [c for c in REQUIRED_PROJECT_COLUMNS if c not in columns]
    if missing:
        raise ValueError(
            f"Predefined projects file '{path}' is missing required "
            f"column(s): {', '.join(missing)}.\n"
            f"Columns found: {', '.join(str(c) for c in columns)}"
        )

    projects = []

    for excel_row, row in rows:

        if row.get("name") is None or not str(row["name"]).strip():
            raise ValueError(
                f"Predefined projects file '{path}', row {excel_row}: 'name' is empty."
            )

        if row.get("animals") is None or not str(row["animals"]).strip():
            raise ValueError(
                f"Predefined projects file '{path}', row {excel_row} "
                f"('{row['name']}'): 'animals' is empty."
            )

        raw_animals = str(row["animals"]).split(";")

        animals = []
        for a in raw_animals:
            a = a.strip()

            if "OR" in a:
                parts = [p.strip() for p in a.split("OR")]
                a = " OR ".join(parts)

            animals.append(a)

        mandatory = False
        if "mandatory" in row:
            val = str(row["mandatory"]).strip().lower()
            mandatory = val in ("true", "1", "yes")

        projects.append({
            "name": str(row["name"]).strip(),
            "animals": animals,
            "theme": str(row.get("theme") or "").strip(),
            "theme_type": str(row.get("theme_type") or "").strip(),
            "mandatory": mandatory
        })

    return projects
